from __future__ import annotations

import argparse
from pathlib import Path

import pdfplumber
from docx import Document


def extract_docx(source: Path) -> str:
    document = Document(source)
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        blocks.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def extract_pdf(source: Path) -> str:
    pages: list[str] = []
    with pdfplumber.open(source) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            pages.append(f"\n===== PAGE {page_number} =====\n{page.extract_text() or ''}")
    return "\n".join(pages)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    extractor = extract_docx if args.source.suffix.lower() == ".docx" else extract_pdf
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(extractor(args.source), encoding="utf-8")


if __name__ == "__main__":
    main()
